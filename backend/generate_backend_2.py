import os

BASE_DIR = r"C:\Users\sudal\.gemini\antigravity\scratch\innovation-dna\backend"

def write_file(path, content):
    full_path = os.path.join(BASE_DIR, path)
    os.makedirs(os.path.dirname(full_path), exist_ok=True)
    with open(full_path, "w", encoding="utf-8") as f:
        f.write(content.strip() + "\n")

FILES = {}

FILES["app/api/evidence.py"] = """
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List, Optional
from app.database import get_db
from app.models.models import User, Evidence, Source
from app.dependencies import get_current_user
from app.schemas.evidence import EvidenceResponse, EvidenceUpdateRequest

router = APIRouter(prefix="/api/evidence", tags=["evidence"])

@router.get("", response_model=List[EvidenceResponse])
def list_evidence(verification_status: Optional[str] = None, source_id: Optional[int] = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    q = db.query(Evidence)
    if verification_status:
        q = q.filter(Evidence.verification_status == verification_status)
    if source_id:
        q = q.filter(Evidence.source_id == source_id)
    items = q.all()
    for item in items:
        if item.source_id:
            item.source = db.query(Source).filter(Source.id == item.source_id).first()
    return items

@router.get("/{id}", response_model=EvidenceResponse)
def get_evidence(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    evidence = db.query(Evidence).filter(Evidence.id == id).first()
    if not evidence:
        raise HTTPException(status_code=404, detail="Evidence not found")
    if evidence.source_id:
        evidence.source = db.query(Source).filter(Source.id == evidence.source_id).first()
    return evidence

@router.put("/{id}", response_model=EvidenceResponse)
def update_evidence(id: int, req: EvidenceUpdateRequest, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    evidence = db.query(Evidence).filter(Evidence.id == id).first()
    if not evidence:
        raise HTTPException(status_code=404, detail="Evidence not found")
    evidence.verification_status = req.verification_status
    db.commit()
    db.refresh(evidence)
    if evidence.source_id:
        evidence.source = db.query(Source).filter(Source.id == evidence.source_id).first()
    return evidence
"""

FILES["app/api/problems.py"] = """
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List, Optional
from app.database import get_db
from app.models.models import User, Problem
from app.dependencies import get_current_user
from app.schemas.problems import ProblemCreate, ProblemUpdate, ProblemResponse

router = APIRouter(prefix="/api/problems", tags=["problems"])

@router.post("", response_model=ProblemResponse)
def create_problem(req: ProblemCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    problem = Problem(**req.model_dump())
    db.add(problem)
    db.commit()
    db.refresh(problem)
    return problem

@router.get("", response_model=List[ProblemResponse])
def list_problems(domain: Optional[str] = None, search: Optional[str] = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    q = db.query(Problem)
    if domain:
        q = q.filter(Problem.domain == domain)
    if search:
        q = q.filter(Problem.title.ilike(f"%{search}%"))
    return q.all()

@router.get("/{id}", response_model=ProblemResponse)
def get_problem(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    problem = db.query(Problem).filter(Problem.id == id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="Problem not found")
    return problem

@router.put("/{id}", response_model=ProblemResponse)
def update_problem(id: int, req: ProblemUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    problem = db.query(Problem).filter(Problem.id == id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="Problem not found")
    for key, val in req.model_dump(exclude_unset=True).items():
        setattr(problem, key, val)
    db.commit()
    db.refresh(problem)
    return problem
"""

FILES["app/api/opportunities.py"] = """
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List
from app.database import get_db
from app.models.models import User, Opportunity, Project, Problem, Evidence, OpportunityEvidence, Experiment
from app.dependencies import get_current_user
from app.schemas.opportunities import OpportunityResponse, OpportunityDetailResponse
from app.schemas.experiments import ExperimentResponse

router = APIRouter(prefix="/api/opportunities", tags=["opportunities"])

@router.get("", response_model=List[OpportunityResponse])
def list_opportunities(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    opps = db.query(Opportunity).all()
    for o in opps:
        p = db.query(Project).filter(Project.id == o.project_id).first()
        pr = db.query(Problem).filter(Problem.id == o.problem_id).first()
        o.project_name = p.name if p else ""
        o.problem_title = pr.title if pr else ""
        o.evidence_count = db.query(OpportunityEvidence).filter(OpportunityEvidence.opportunity_id == o.id).count()
    return opps

@router.get("/{id}", response_model=OpportunityResponse)
def get_opportunity(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    o = db.query(Opportunity).filter(Opportunity.id == id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Opportunity not found")
    p = db.query(Project).filter(Project.id == o.project_id).first()
    pr = db.query(Problem).filter(Problem.id == o.problem_id).first()
    o.project_name = p.name if p else ""
    o.problem_title = pr.title if pr else ""
    o.evidence_count = db.query(OpportunityEvidence).filter(OpportunityEvidence.opportunity_id == o.id).count()
    return o

@router.get("/{id}/evidence")
def get_evidence_for_opp(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    o = db.query(Opportunity).filter(Opportunity.id == id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Opportunity not found")
    links = db.query(OpportunityEvidence).filter(OpportunityEvidence.opportunity_id == id).all()
    evidence_ids = [l.evidence_id for l in links]
    return db.query(Evidence).filter(Evidence.id.in_(evidence_ids)).all()

@router.post("/{id}/experiment", response_model=ExperimentResponse)
def create_experiment(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    o = db.query(Opportunity).filter(Opportunity.id == id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Opportunity not found")
    e = Experiment(opportunity_id=id, title=f"Experiment for {o.title}", methodology="A/B Testing", success_criteria="Statistically significant improvement", status="planned")
    db.add(e)
    db.commit()
    db.refresh(e)
    return e

@router.get("/{id}/experiment", response_model=List[ExperimentResponse])
def get_experiments(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Experiment).filter(Experiment.opportunity_id == id).all()
"""

FILES["app/api/experiments.py"] = """
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List
from app.database import get_db
from app.models.models import User, Experiment, ExperimentResult
from app.dependencies import get_current_user
from app.schemas.experiments import ExperimentResponse, ExperimentUpdate, ExperimentResultCreate, ExperimentResultResponse

router = APIRouter(prefix="/api/experiments", tags=["experiments"])

@router.get("/{id}", response_model=ExperimentResponse)
def get_experiment(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    e = db.query(Experiment).filter(Experiment.id == id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Experiment not found")
    return e

@router.put("/{id}", response_model=ExperimentResponse)
def update_experiment(id: int, req: ExperimentUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    e = db.query(Experiment).filter(Experiment.id == id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Experiment not found")
    for key, val in req.model_dump(exclude_unset=True).items():
        setattr(e, key, val)
    db.commit()
    db.refresh(e)
    return e

@router.post("/{id}/results", response_model=ExperimentResultResponse)
def add_result(id: int, req: ExperimentResultCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    r = ExperimentResult(**req.model_dump(), experiment_id=id)
    db.add(r)
    db.commit()
    db.refresh(r)
    return r

@router.get("/{id}/results", response_model=List[ExperimentResultResponse])
def get_results(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(ExperimentResult).filter(ExperimentResult.experiment_id == id).all()
"""

FILES["app/api/jobs.py"] = """
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List
from app.database import get_db
from app.models.models import User, Job
from app.dependencies import get_current_user
from app.schemas.jobs import JobResponse

router = APIRouter(prefix="/api/jobs", tags=["jobs"])

@router.get("/{id}", response_model=JobResponse)
def get_job(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    job = db.query(Job).filter(Job.id == id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return job

@router.get("", response_model=List[JobResponse])
def list_jobs(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    # Simple returning all for now, or filter by user projects
    return db.query(Job).all()
"""

FILES["app/api/search.py"] = """
from fastapi import APIRouter, Depends
from sqlalchemy.orm import Session
from typing import List, Optional
from app.database import get_db
from app.models.models import User, Project, Problem, Source, Opportunity
from app.dependencies import get_current_user
from app.schemas.search import SearchResponse, SearchResultItem

router = APIRouter(prefix="/api/search", tags=["search"])

@router.get("", response_model=SearchResponse)
def search(q: str, type: Optional[str] = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    results = []
    if not type or type == "project":
        ps = db.query(Project).filter(Project.name.ilike(f"%{q}%")).all()
        for p in ps:
            results.append(SearchResultItem(type="project", id=p.id, title=p.name, subtitle=p.domain))
    if not type or type == "problem":
        pr = db.query(Problem).filter(Problem.title.ilike(f"%{q}%") | Problem.domain.ilike(f"%{q}%")).all()
        for p in pr:
            results.append(SearchResultItem(type="problem", id=p.id, title=p.title, subtitle=p.domain))
    if not type or type == "source":
        ss = db.query(Source).filter(Source.title.ilike(f"%{q}%")).all()
        for s in ss:
            results.append(SearchResultItem(type="source", id=s.id, title=s.title, subtitle=s.publisher))
    if not type or type == "opportunity":
        ops = db.query(Opportunity).filter(Opportunity.title.ilike(f"%{q}%")).all()
        for o in ops:
            results.append(SearchResultItem(type="opportunity", id=o.id, title=o.title, subtitle=o.status))
            
    return SearchResponse(results=results)
"""

FILES["app/api/admin.py"] = """
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.models import User, Source, Project, Problem, Evidence, Job, AnalysisRun
from app.dependencies import require_admin

router = APIRouter(prefix="/api/admin", tags=["admin"])

@router.get("/stats")
def stats(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return {
        "sources": db.query(Source).count(),
        "projects": db.query(Project).count(),
        "problems": db.query(Problem).count(),
        "evidence": db.query(Evidence).count(),
        "jobs": db.query(Job).count(),
        "users": db.query(User).count()
    }

@router.get("/sources")
def get_sources(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return db.query(Source).all()

@router.get("/projects")
def get_projects(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return db.query(Project).all()

@router.get("/problems")
def get_problems(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return db.query(Problem).all()

@router.get("/evidence")
def get_evidence(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return db.query(Evidence).all()

@router.get("/jobs")
def get_jobs(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return db.query(Job).all()

@router.get("/analysis-runs")
def get_analysis_runs(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    return db.query(AnalysisRun).all()

@router.post("/sources/{id}/verify")
def verify_source(id: int, db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    s = db.query(Source).filter(Source.id == id).first()
    if s:
        s.verification_status = "verified"
        db.commit()
    return s

@router.post("/evidence/{id}/verify")
def verify_evidence(id: int, db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    e = db.query(Evidence).filter(Evidence.id == id).first()
    if e:
        e.verification_status = "verified"
        db.commit()
    return e

@router.post("/jobs/{id}/retry")
def retry_job(id: int, db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    j = db.query(Job).filter(Job.id == id).first()
    if j:
        j.status = "queued"
        db.commit()
    return j
"""

FILES["app/services/__init__.py"] = ""

FILES["app/services/storage_service.py"] = """
import os
import uuid
import hashlib
from pathlib import Path

class LocalStorageService:
    def __init__(self, base_path: str = "./uploads"):
        self.base_path = Path(base_path)
        self.base_path.mkdir(parents=True, exist_ok=True)
    
    def save(self, content: bytes, original_filename: str) -> tuple[str, str]:
        ext = Path(original_filename).suffix
        filename = f"{uuid.uuid4().hex}{ext}"
        filepath = self.base_path / filename
        filepath.write_bytes(content)
        checksum = hashlib.md5(content).hexdigest()
        return str(filepath), checksum
    
    def get(self, path: str) -> bytes:
        return Path(path).read_bytes()
    
    def delete(self, path: str) -> None:
        p = Path(path)
        if p.exists():
            p.unlink()
"""

FILES["app/ingestion/__init__.py"] = ""

FILES["app/ingestion/parser.py"] = """
import io
from pathlib import Path

def parse_document(filepath: str, mime_type: str) -> str:
    \"\"\"Parse document and return extracted text.\"\"\"
    filepath = Path(filepath)
    
    if mime_type == "application/pdf" or filepath.suffix == ".pdf":
        return parse_pdf(filepath)
    elif mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or filepath.suffix == ".docx":
        return parse_docx(filepath)
    elif mime_type in ("application/vnd.openxmlformats-officedocument.presentationml.presentation",) or filepath.suffix == ".pptx":
        return parse_pptx(filepath)
    elif filepath.suffix in (".txt", ".md"):
        return filepath.read_text(encoding="utf-8", errors="ignore")
    else:
        return filepath.read_text(encoding="utf-8", errors="ignore")

def parse_pdf(filepath):
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(str(filepath)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        return "\\n\\n".join(text_parts)
    except Exception as e:
        return f"[PDF parsing error: {e}]"

def parse_docx(filepath):
    try:
        from docx import Document
        doc = Document(str(filepath))
        return "\\n\\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[DOCX parsing error: {e}]"

def parse_pptx(filepath):
    try:
        from pptx import Presentation
        prs = Presentation(str(filepath))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
        return "\\n\\n".join(texts)
    except Exception as e:
        return f"[PPTX parsing error: {e}]"
"""

FILES["app/ingestion/chunker.py"] = """
def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        i += chunk_size - overlap
    return chunks if chunks else [text[:2000]] if text else []
"""

for path, content in FILES.items():
    write_file(path, content)

print("Batch 2 created successfully.")
