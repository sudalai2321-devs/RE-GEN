import io
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from typing import List
from app.database import get_db
from app.models.models import User, Project, Document
from app.dependencies import get_current_user
from app.schemas.documents import DocumentResponse
from app.services.storage_service import LocalStorageService

router = APIRouter(prefix="/api/projects", tags=["documents"])
storage = LocalStorageService()

@router.post("/{id}/documents", response_model=DocumentResponse)
async def upload_document(id: int, file: UploadFile = File(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project = db.query(Project).filter(Project.id == id, Project.user_id == current_user.id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
        
    content = await file.read()
    filepath, checksum = storage.save(content, file.filename)
    
    extracted_text = ""
    try:
        if file.filename.lower().endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            extracted_pages = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    extracted_pages.append(f"--- Page {i+1} ---\n{text.strip()}")
            extracted_text = "\n\n".join(extracted_pages)
        elif file.filename.lower().endswith(".pptx"):
            from pptx import Presentation
            prs = Presentation(io.BytesIO(content))
            slide_texts = []
            for i, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t:
                                texts.append(t)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_texts:
                                texts.append(" | ".join(row_texts))
                if texts:
                    slide_texts.append(f"--- Slide {i+1} ---\n" + "\n".join(texts))
            extracted_text = "\n\n".join(slide_texts)
        elif file.filename.lower().endswith(".docx"):
            try:
                from docx import Document as DocxDocument
                doc_obj = DocxDocument(io.BytesIO(content))
                paragraphs = [p.text.strip() for p in doc_obj.paragraphs if p.text.strip()]
                extracted_text = "\n\n".join(paragraphs)
            except ImportError:
                extracted_text = ""
        elif file.filename.lower().endswith((".txt", ".md", ".json")):
            extracted_text = content.decode("utf-8", errors="ignore")
    except Exception as err:
        import traceback
        print(f"Warning: Text extraction failed for {file.filename}: {err}")
        traceback.print_exc()

    # Intelligent inspection of extracted document text
    is_blank_template = False
    if extracted_text:
        lower_txt = extracted_text.lower()
        template_indicators = [
            "state the problem you want to tackle",
            "state the solution of your problem statement",
            "tech stack used:\n\n--- slide",
            "tech stack used:\n--- slide",
            "tech stack used:\n\n\n",
        ]
        if any(ind in lower_txt for ind in template_indicators):
            is_blank_template = True
        elif "hackathon" in lower_txt and len(extracted_text.strip()) < 180 and "team name" in lower_txt:
            is_blank_template = True

    doc = Document(
        project_id=project.id,
        filename=file.filename,
        original_filename=file.filename,
        storage_path=filepath,
        file_path=filepath,
        mime_type=file.content_type or "application/octet-stream",
        file_size=len(content),
        size_bytes=len(content),
        checksum=checksum,
        content_hash=checksum,
        processing_status="template_detected" if is_blank_template else "completed",
        status="uploaded",
        extracted_text=extracted_text[:20000] if extracted_text else ""
    )
    db.add(doc)

    # Enrich project metadata intelligently
    if is_blank_template:
        project.objective = f"Hackathon Pitch Deck Template ({file.filename}) — Celestia / ISTE VIT Vellore track submission."
        project.problem = "Unfilled hackathon template: Problem statement, technical solution, and tech stack are currently unpopulated."
        project.failure_summary = "Template document uploaded without technical body. No system failure or boundary limitation documented yet."
        if not project.domain or project.domain == "Industrial IoT":
            project.domain = "Open Innovation"
    elif extracted_text:
        # Check for structured hackathon sections (like in TechTitans-RE-Gen.pptx)
        lower_txt = extracted_text.lower()
        if "problem statement:" in lower_txt and "proposed  solution:" in lower_txt:
            try:
                # Extract problem statement
                prob_part = extracted_text.split("PROBLEM STATEMENT:", 1)[1]
                prob_text = prob_part.split("--- Slide", 1)[0].split("PROPOSED", 1)[0].strip()
                if prob_text:
                    project.problem = prob_text[:500]
                
                # Extract proposed solution
                sol_part = extracted_text.split("PROPOSED", 1)[1]
                sol_part_clean = sol_part.split("SOLUTION:", 1)[-1] if "SOLUTION:" in sol_part else sol_part
                sol_text = sol_part_clean.split("--- Slide", 1)[0].split("TECH STACK", 1)[0].strip()
                if sol_text:
                    project.objective = sol_text[:500]
                
                # Auto-infer domain
                if "re:gen" in lower_txt or "innovation" in lower_txt or "open innovation" in lower_txt:
                    project.domain = "AI & Open Innovation"
            except Exception:
                pass
        else:
            snippet = extracted_text[:500].replace("\n", " ").strip()
            if not project.problem or "Documented technical attempt" in project.problem or "Identify innovation gaps" in project.problem:
                project.problem = f"Extracted from {file.filename}: {snippet[:300]}..."
            if not project.objective or "Ingested from document" in project.objective or "Analyze technical capabilities" in project.objective:
                project.objective = f"Analysis of technical capabilities and constraints documented in {file.filename}."
    db.commit()
    db.refresh(doc)
    return doc

@router.get("/{id}/documents", response_model=List[DocumentResponse])
def list_documents(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project = db.query(Project).filter(Project.id == id, Project.user_id == current_user.id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    docs = db.query(Document).filter(Document.project_id == id).all()
    return docs
