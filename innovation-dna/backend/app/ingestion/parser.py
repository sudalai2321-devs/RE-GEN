import io
from pathlib import Path

def parse_document(filepath: str, mime_type: str) -> str:
    """Parse document and return extracted text."""
    filepath = Path(filepath)
    
    if mime_type == "application/pdf" or filepath.suffix == ".pdf":
        return parse_pdf(filepath)
    elif mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or filepath.suffix == ".docx":
        return parse_docx(filepath)
    elif mime_type in ("application/vnd.openxmlformats-officedocument.presentationml.presentation",) or filepath.suffix == ".pptx":
        return parse_pptx(filepath)
    elif filepath.suffix in (".txt", ".md"):
        return filepath.read_text(encoding="utf-8", errors="ignore")
    else:
        return filepath.read_text(encoding="utf-8", errors="ignore")

def parse_pdf(filepath):
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(str(filepath)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        return "\n\n".join(text_parts)
    except Exception as e:
        return f"[PDF parsing error: {e}]"

def parse_docx(filepath):
    try:
        from docx import Document
        doc = Document(str(filepath))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[DOCX parsing error: {e}]"

def parse_pptx(filepath):
    try:
        from pptx import Presentation
        prs = Presentation(str(filepath))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
        return "\n\n".join(texts)
    except Exception as e:
        return f"[PPTX parsing error: {e}]"
